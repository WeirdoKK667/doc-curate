"""
build_toc_docx.py · 阶段 5 目录 docx

生成独立小 docx（仅章节标题 + PDF 页标记）。
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.shared import RGBColor


def extract_toc(process_dir: Path) -> list:
    """从 zh.md 抽取目录（标题 + 页码）。"""
    toc = []
    for md_file in sorted(process_dir.glob("p*.zh.md")):
        text = md_file.read_text(encoding="utf-8")
        page_match = re.search(r"<!-- page:(\d+) -->", text)
        page_num = int(page_match.group(1)) if page_match else 0

        # 提取 H1-H3 标题
        for match in re.finditer(r"^(#{1,3})\s+(.+)$", text, re.MULTILINE):
            level = len(match.group(1))
            title = match.group(2).strip()
            toc.append({"level": level, "title": title, "page": page_num})
    return toc


def build_toc_docx(toc: list, output_path: Path) -> None:
    """生成目录 docx。"""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.element.rPr.rFonts.eastAsia = "宋体"

    title = doc.add_heading("目录", level=0)
    title.runs[0].font.size = Pt(18)
    title.runs[0].font.element.rPr.rFonts.eastAsia = "黑体"

    for entry in toc:
        p = doc.add_paragraph()
        run = p.add_run(f"{entry['title']}    PDF 第 {entry['page']} 页")
        run.font.size = Pt(11 + (3 - entry["level"]) * 2)
        run.font.element.rPr.rFonts.eastAsia = "宋体"
        p.paragraph_format.left_indent = Pt((entry["level"] - 1) * 20)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"已生成目录: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="阶段 5 目录 docx")
    parser.add_argument("--process-dir", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    toc = extract_toc(Path(args.process_dir))
    build_toc_docx(toc, Path(args.output))


if __name__ == "__main__":
    main()
